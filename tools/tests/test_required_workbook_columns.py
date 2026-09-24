"""A workbook column the generator reads and a sheet lacks is refused, not read
as empty. Issue #1882.

WHAT WAS WRONG. `tools/generate_datatables.py` reads the design workbook by
column name, and `_cell` returned an empty string for a column the sheet did not
have -- the same as for an empty cell. So a column deleted or renamed in the
workbook emptied that field on every row, and nothing said so.

WHAT HOLDS IT NOW. `_cell` raises `DataError`, naming the sheet and the column,
unless `OPTIONAL_COLUMNS` declares that column optional for that sheet. Ruled on
2026-09-23 under the owner's delegation: every column the generator asks for is
required, and optional only by declaration.

THE REAL WORKBOOK IS CHECKED ELSEWHERE, BY RUNNING THE GENERATOR.
`test_generate_datatables.py::...::test_the_committed_csvs_are_current` runs
`generate_datatables.main(["--check"])` over `docs/All_Things_Cataclysm.xlsx`,
so every column read on the way is held to this rule there. This file holds the
rule itself, on hand-made headers, and holds `OPTIONAL_COLUMNS` to the workbook
so that it only shrinks.
"""

from __future__ import annotations

import pathlib
import sys

import openpyxl
import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import generate_datatables as gen  # noqa: E402

WORKBOOK = gen.REPO_ROOT / "docs" / "All_Things_Cataclysm.xlsx"

ROWS = [["Name", "Stat", "Kind"],
        ["A", "armor", ""]]


def headers():
    return gen._header_index(ROWS, "Hand Made")


def test_a_column_the_sheet_lacks_is_refused_and_named():
    with pytest.raises(gen.DataError) as refused:
        gen._cell(ROWS[1], headers(), "Scale")
    message = str(refused.value)
    assert "Hand Made" in message and "'Scale'" in message, message


def test_a_column_the_sheet_has_and_leaves_empty_is_still_empty():
    """The control: an empty cell in a column that exists is not refused."""
    assert gen._cell(ROWS[1], headers(), "Kind") == ""
    assert gen._cell(ROWS[1], headers(), "Stat") == "armor"


def test_a_row_shorter_than_its_header_reads_its_last_columns_as_empty():
    """A spreadsheet drops trailing empty cells from a row; that is not a
    missing column."""
    assert gen._cell(["A"], headers(), "Kind") == ""


def test_a_declared_optional_column_reads_as_empty(monkeypatch):
    monkeypatch.setattr(gen, "OPTIONAL_COLUMNS",
                        {"Hand Made": {"Scale": "issue #0000, a test"}})
    assert gen._cell(ROWS[1], headers(), "Scale") == ""


def test_an_optional_column_on_another_sheet_does_not_excuse_this_one(monkeypatch):
    monkeypatch.setattr(gen, "OPTIONAL_COLUMNS",
                        {"Some Other Sheet": {"Scale": "issue #0000, a test"}})
    with pytest.raises(gen.DataError):
        gen._cell(ROWS[1], headers(), "Scale")


def stale_optional_columns(book, table: dict[str, dict[str, str]]) -> list[str]:
    """Every declared optional column whose sheet now has it, as 'sheet: column'."""
    stale = []
    for sheet, columns in table.items():
        assert sheet in book.sheetnames, (
            f"OPTIONAL_COLUMNS names the sheet {sheet!r}, which the workbook "
            f"does not have")
        first = next(book[sheet].iter_rows(values_only=True, max_row=1), ())
        present = {gen.clean(value) for value in first if gen.clean(value)}
        stale += [f"{sheet}: {column}" for column in columns if column in present]
    return stale


def test_an_optional_column_the_sheet_now_has_is_not_listed():
    """What makes OPTIONAL_COLUMNS shrink rather than rot, on the real workbook."""
    book = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    assert len(book.sheetnames) >= 10, (
        f"only {len(book.sheetnames)} sheets were read from {WORKBOOK}")
    stale = stale_optional_columns(book, gen.OPTIONAL_COLUMNS)
    assert not stale, (
        "these columns are declared optional in OPTIONAL_COLUMNS in "
        "tools/generate_datatables.py, and the workbook now has them:\n"
        + "\n".join(f"  {entry}" for entry in stale)
        + "\n\nDelete each entry in the same change that added the column, so "
        "the column is required from then on. Issue #1882.")


def test_the_shrink_check_fires_on_a_listed_column_that_exists(tmp_path):
    """The check above, on a hand-made workbook, so it is shown to act while
    OPTIONAL_COLUMNS is empty."""
    book = openpyxl.Workbook()
    book.active.title = "Hand Made"
    book.active.append(ROWS[0])
    path = tmp_path / "w.xlsx"
    book.save(path)
    loaded = openpyxl.load_workbook(path, read_only=True)

    table = {"Hand Made": {"Stat": "issue #0000", "Scale": "issue #0000"}}
    assert stale_optional_columns(loaded, table) == ["Hand Made: Stat"]
